"""File extraction and token-aware processing (no downloads at import)."""
from __future__ import annotations

import io
import re
from pathlib import Path

LABEL_NAMES = ("negative", "neutral", "positive")
MAX_FILE_BYTES = 10 * 1024 * 1024


def clean_text(text: str) -> str:
    text = text.replace("\x00", " ").replace("\ufeff", "")
    text = re.sub(r"\r\n?", "\n", text)
    text = re.sub(r"[ \t]+", " ", text)
    return re.sub(r"\n{3,}", "\n\n", text).strip()


def extract_file(name: str, data: bytes) -> str:
    if len(data) > MAX_FILE_BYTES:
        raise ValueError("This file is larger than 10 MB. Upload the news document, not the model ZIP.")
    extension = Path(name).suffix.lower()
    if extension == ".pdf":
        from pypdf import PdfReader
        reader = PdfReader(io.BytesIO(data))
        if reader.is_encrypted and not reader.decrypt(""):
            raise ValueError("This PDF is password-protected. Please upload an unencrypted copy.")
        if len(reader.pages) > 30:
            raise ValueError("Please upload a PDF of at most 30 pages containing the relevant article.")
        text = "\n\n".join(page.extract_text() or "" for page in reader.pages)
    elif extension == ".docx":
        from docx import Document
        document = Document(io.BytesIO(data))
        paragraphs = [p.text for p in document.paragraphs if p.text.strip()]
        for table in document.tables:
            for row in table.rows:
                paragraphs.append(" | ".join(cell.text for cell in row.cells))
        text = "\n\n".join(paragraphs)
    elif extension in {".txt", ".md"}:
        for encoding in ("utf-8-sig", "utf-16", "gb18030"):
            try:
                text = data.decode(encoding)
                break
            except UnicodeError:
                continue
        else:
            raise ValueError("The text encoding could not be read. Please save the file as UTF-8 TXT.")
    else:
        raise ValueError("Supported formats: PDF, TXT, Markdown and DOCX.")
    text = clean_text(text)
    if not text:
        raise ValueError("No text could be extracted. Scanned PDFs need OCR before uploading.")
    return text


def token_chunks(text: str, tokenizer, max_length: int = 128, max_chunks: int = 32):
    ids = tokenizer.encode(text, add_special_tokens=False, truncation=False)
    payload = max_length - tokenizer.num_special_tokens_to_add(pair=False)
    if payload < 1:
        raise ValueError("max_length is smaller than the number of required special tokens.")
    groups = [ids[i:i + payload] for i in range(0, len(ids), payload)]
    processed = groups[:max_chunks]
    texts = [tokenizer.decode(g, skip_special_tokens=True) for g in processed]
    return texts, {
        "tokens": len(ids), "used_tokens": sum(len(g) for g in processed),
        "chunks": len(processed), "total_chunks": len(groups),
        "truncated": len(groups) > max_chunks,
    }


def probability_rows(raw):
    if raw and isinstance(raw[0], list):
        raw = raw[0]
    rows = []
    for item in raw:
        label = str(item["label"]).strip().lower()
        if label not in LABEL_NAMES:
            raise ValueError(
                f"Unrecognized model label: {label!r}. config.json must specify "
                "negative / neutral / positive, not ambiguous LABEL_0 values."
            )
        rows.append({"label": label, "score": float(item["score"])})
    if {r["label"] for r in rows} != set(LABEL_NAMES):
        raise ValueError("The model must output negative, neutral and positive labels.")
    return rows
